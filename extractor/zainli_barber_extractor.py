import json, threading, time, webbrowser
from urllib import request, error
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
except Exception:
    Workbook = None

PLACES_URL = 'https://places.googleapis.com/v1/places:searchText'
SUPABASE_URL = 'https://lxdqwwdhjgxnhysxlodi.supabase.co'
SUPABASE_KEY = 'sb_publishable_hfX5zq7_74L1esoHkBqlhg_b7mVr0Hg'
BAGHDAD_AREAS = ['المنصور','الكرادة','الجادرية','زيونة','بغداد الجديدة','الأعظمية','الكاظمية','السيدية','الدورة','العامرية','الغزالية','اليرموك','الحارثية','الشعب','البلديات','شارع فلسطين','البياع','حي الجامعة','الزعفرانية','حي العامل','الحرية','الصالحية','الوشاش']
FIELD_MASK = ','.join([
    'places.id','places.displayName','places.formattedAddress',
    'places.nationalPhoneNumber','places.internationalPhoneNumber',
    'places.location','places.rating','places.userRatingCount',
    'places.regularOpeningHours','places.googleMapsUri','places.websiteUri',
    'places.businessStatus','nextPageToken'
])

def http_json(url, method='GET', headers=None, body=None, timeout=35):
    data = None if body is None else json.dumps(body, ensure_ascii=False).encode('utf-8')
    req = request.Request(url, data=data, method=method, headers=headers or {})
    try:
        with request.urlopen(req, timeout=timeout) as r:
            raw = r.read().decode('utf-8')
            return json.loads(raw) if raw else None
    except error.HTTPError as e:
        raw = e.read().decode('utf-8', errors='replace')
        raise RuntimeError(f'HTTP {e.code}: {raw[:500]}')
    except Exception as e:
        raise RuntimeError(str(e))

def place_name(p):
    d = p.get('displayName') or {}
    return d.get('text') or ''

def phone_of(p):
    return p.get('nationalPhoneNumber') or p.get('internationalPhoneNumber') or ''

def hours_of(p):
    return ' | '.join((p.get('regularOpeningHours') or {}).get('weekdayDescriptions') or [])

def normalize_phone(s):
    s = str(s or '').strip()
    for i,ch in enumerate('٠١٢٣٤٥٦٧٨٩'):
        s = s.replace(ch, str(i))
    return ''.join(c for c in s if c.isdigit() or c == '+')

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title('زينلي - استخراج الحلاقين من Google Maps')
        self.geometry('1450x820')
        self.minsize(1050,650)
        self.rows=[]
        self.stop_event=threading.Event()
        self.build_ui()

    def build_ui(self):
        top=ttk.Frame(self,padding=12); top.pack(fill='x')
        ttk.Label(top,text='Google Places API Key').grid(row=0,column=0,sticky='w')
        self.api=ttk.Entry(top,width=50,show='•'); self.api.grid(row=1,column=0,columnspan=2,sticky='ew',padx=(0,8))
        ttk.Label(top,text='المدينة').grid(row=0,column=2,sticky='w')
        self.city=ttk.Entry(top,width=22); self.city.insert(0,'Baghdad, Iraq'); self.city.grid(row=1,column=2,sticky='ew',padx=(0,8))
        ttk.Label(top,text='كلمات البحث - افصل بـ ;').grid(row=0,column=3,sticky='w')
        self.terms=ttk.Entry(top,width=35); self.terms.insert(0,'حلاق;صالون رجالي;barber'); self.terms.grid(row=1,column=3,sticky='ew',padx=(0,8))
        self.use_areas=tk.BooleanVar(value=True); ttk.Checkbutton(top,text='قسّم بغداد لمناطق حتى يجيب نتائج أكثر',variable=self.use_areas).grid(row=2,column=0,columnspan=2,sticky='w',pady=(8,0))
        self.only_phone=tk.BooleanVar(value=True); ttk.Checkbutton(top,text='فقط اللي عندهم رقم هاتف',variable=self.only_phone).grid(row=2,column=2,sticky='w',pady=(8,0))
        ttk.Button(top,text='🔍 بحث',command=self.start_search).grid(row=1,column=4,padx=4)
        ttk.Button(top,text='⛔ إيقاف',command=self.stop_search).grid(row=2,column=4,padx=4,pady=(8,0))
        for c in range(4): top.columnconfigure(c,weight=1)

        ttk.Label(self,text='يحتاج Google Places API (New). رقم الهاتف ليس متوفراً لكل محل. البحث النصي محدود تقريباً إلى 60 نتيجة لكل استعلام، لذلك البرنامج يقسم بغداد ويزيل التكرار.',padding=(12,0,12,8),foreground='#555').pack(fill='x')

        bar=ttk.Frame(self,padding=(12,4)); bar.pack(fill='x')
        ttk.Button(bar,text='📊 تصدير Excel',command=self.export_excel).pack(side='left',padx=4)
        ttk.Button(bar,text='🌐 فتح Google Maps',command=self.open_map).pack(side='left',padx=4)
        ttk.Button(bar,text='⬆ رفع المحدد إلى Zainli',command=self.upload_selected).pack(side='left',padx=4)
        ttk.Button(bar,text='⬆ رفع الكل إلى Zainli',command=self.upload_all).pack(side='left',padx=4)
        ttk.Button(bar,text='🗑 مسح',command=self.clear).pack(side='left',padx=4)
        ttk.Label(bar,text='Admin Code:').pack(side='right',padx=(8,4))
        self.admin=ttk.Entry(bar,width=18,show='•'); self.admin.pack(side='right')

        cols=('name','phone','address','area','rating','reviews','lat','lng','maps','website','status')
        self.tree=ttk.Treeview(self,columns=cols,show='headings',selectmode='extended')
        labels={'name':'الاسم','phone':'الهاتف','address':'العنوان','area':'منطقة البحث','rating':'التقييم','reviews':'التقييمات','lat':'Lat','lng':'Lng','maps':'Google Maps','website':'Website','status':'الحالة'}
        widths={'name':190,'phone':130,'address':260,'area':120,'rating':70,'reviews':80,'lat':85,'lng':85,'maps':180,'website':160,'status':100}
        for c in cols:
            self.tree.heading(c,text=labels[c]); self.tree.column(c,width=widths[c],anchor='w')
        self.tree.pack(fill='both',expand=True,padx=12,pady=6)
        self.tree.bind('<Double-1>',lambda e:self.open_map())

        bottom=ttk.Frame(self,padding=12); bottom.pack(fill='x')
        self.prog=ttk.Progressbar(bottom,mode='indeterminate'); self.prog.pack(side='left',fill='x',expand=True,padx=(0,10))
        self.status=ttk.Label(bottom,text='جاهز'); self.status.pack(side='right')

    def set_status(self,t): self.after(0,lambda:self.status.config(text=t))
    def start_search(self):
        if not self.api.get().strip(): messagebox.showwarning('API Key','حط Google Places API Key أولاً.'); return
        self.stop_event.clear(); self.prog.start(10); threading.Thread(target=self.search_worker,daemon=True).start()
    def stop_search(self): self.stop_event.set(); self.set_status('جاري الإيقاف...')

    def search_worker(self):
        try:
            key=self.api.get().strip(); city=self.city.get().strip(); terms=[x.strip() for x in self.terms.get().split(';') if x.strip()]
            areas=BAGHDAD_AREAS if self.use_areas.get() and 'baghdad' in city.lower() else ['']
            queries=[f'{term} {area} {city}'.strip() for area in areas for term in terms]
            known={r['place_id'] for r in self.rows}
            added=0
            for qi,q in enumerate(queries,1):
                if self.stop_event.is_set(): break
                self.set_status(f'بحث {qi}/{len(queries)}: {q}')
                token=None; got=0
                while True:
                    if self.stop_event.is_set(): break
                    body={'textQuery':q,'pageSize':20,'languageCode':'ar','regionCode':'IQ'}
                    if token: body['pageToken']=token
                    data=http_json(PLACES_URL,'POST',{'Content-Type':'application/json','X-Goog-Api-Key':key,'X-Goog-FieldMask':FIELD_MASK},body)
                    for p in (data or {}).get('places') or []:
                        pid=p.get('id') or p.get('name')
                        if not pid or pid in known: continue
                        ph=phone_of(p)
                        if self.only_phone.get() and not ph: continue
                        loc=p.get('location') or {}
                        row={'place_id':pid,'name':place_name(p),'phone':ph,'address':p.get('formattedAddress') or '',
                             'area':self.area_from_query(q),'rating':p.get('rating') or '','reviews':p.get('userRatingCount') or '',
                             'lat':loc.get('latitude') or '','lng':loc.get('longitude') or '','maps':p.get('googleMapsUri') or '',
                             'website':p.get('websiteUri') or '','status':p.get('businessStatus') or '',
                             'hours':hours_of(p),'search_query':q}
                        known.add(pid); self.rows.append(row); got+=1; added+=1
                        self.after(0,lambda r=row:self.insert(r))
                    token=(data or {}).get('nextPageToken')
                    if not token or got>=60: break
                    time.sleep(.3)
            self.set_status(f'تم - أضيف {added} | الإجمالي {len(self.rows)}')
        except Exception as e:
            self.after(0,lambda:messagebox.showerror('خطأ',str(e))); self.set_status('خطأ')
        finally: self.after(0,self.prog.stop)

    def area_from_query(self,q):
        for a in BAGHDAD_AREAS:
            if a in q: return a
        return self.city.get().strip()

    def insert(self,r):
        self.tree.insert('', 'end', values=(r['name'],r['phone'],r['address'],r['area'],r['rating'],r['reviews'],r['lat'],r['lng'],r['maps'],r['website'],r['status']))
    def clear(self):
        self.rows.clear(); [self.tree.delete(x) for x in self.tree.get_children()]; self.set_status('تم المسح')
    def selected_indices(self):
        all_ids=list(self.tree.get_children()); return [all_ids.index(i) for i in self.tree.selection() if i in all_ids]
    def open_map(self):
        idx=self.selected_indices()
        if not idx: messagebox.showinfo('اختيار','حدد صف أولاً.'); return
        url=self.rows[idx[0]].get('maps')
        if url: webbrowser.open(url)

    def export_excel(self):
        if not self.rows: messagebox.showinfo('Excel','ماكو نتائج.'); return
        if Workbook is None: messagebox.showerror('Excel','openpyxl غير مثبت.'); return
        path=filedialog.asksaveasfilename(defaultextension='.xlsx',filetypes=[('Excel','*.xlsx')],initialfile='Zainli_Barbers.xlsx')
        if not path: return
        wb=Workbook(); ws=wb.active; ws.title='Barbers'
        headers=['Name','Phone','Address','Search Area','Rating','Reviews','Latitude','Longitude','Google Maps','Website','Business Status','Opening Hours','Place ID','Search Query']
        ws.append(headers)
        for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F4E78')
        for r in self.rows: ws.append([r['name'],r['phone'],r['address'],r['area'],r['rating'],r['reviews'],r['lat'],r['lng'],r['maps'],r['website'],r['status'],r['hours'],r['place_id'],r['search_query']])
        ws.freeze_panes='A2'; wb.save(path); messagebox.showinfo('تم',f'انحفظ:\n{path}')

    def upload_selected(self):
        idx=self.selected_indices()
        if not idx: messagebox.showinfo('Zainli','حدد صفوف أولاً.'); return
        self.start_upload([self.rows[i] for i in idx])
    def upload_all(self): self.start_upload(list(self.rows))
    def start_upload(self,rows):
        code=self.admin.get().strip()
        if not code: messagebox.showwarning('Admin Code','اكتب Admin Code.'); return
        rows=[r for r in rows if normalize_phone(r.get('phone'))]
        if not rows: messagebox.showinfo('Zainli','ماكو صفوف بيها رقم.'); return
        if not messagebox.askyesno('تأكيد',f'رفع {len(rows)} حلاق إلى Zainli؟'): return
        self.prog.start(10); threading.Thread(target=self.upload_worker,args=(rows,code),daemon=True).start()

    def upload_worker(self,rows,code):
        ok=dup=fail=0
        for i,r in enumerate(rows,1):
            try:
                self.set_status(f'رفع {i}/{len(rows)}: {r["name"]}')
                body={'p_code':code,'p_name':r['name'] or 'Barber','p_shop':r['name'] or 'Barber','p_phone':normalize_phone(r['phone']),
                      'p_area':r['area'] or r['address'][:80] or 'Baghdad','p_lat':float(r['lat'] or 33.31),'p_lng':float(r['lng'] or 44.36)}
                http_json(SUPABASE_URL+'/rest/v1/rpc/z_admin_add_barber','POST',{'Content-Type':'application/json','apikey':SUPABASE_KEY},body,25); ok+=1
            except Exception as e:
                s=str(e).lower()
                if 'duplicate' in s or 'unique' in s: dup+=1
                else: fail+=1
        self.after(0,self.prog.stop); self.set_status(f'Zainli: نجح {ok} | مكرر {dup} | فشل {fail}')
        self.after(0,lambda:messagebox.showinfo('Zainli',f'نجح: {ok}\nمكرر: {dup}\nفشل: {fail}'))

if __name__=='__main__': App().mainloop()
